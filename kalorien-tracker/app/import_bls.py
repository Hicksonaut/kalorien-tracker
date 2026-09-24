"""Import Bundeslebensmittelschluessel (BLS) 4.0 nach `data/bls.db`.

Quelle: Max Rubner-Institut (2025): Bundeslebensmittelschluessel (BLS), Version 4.0 -
Deutsche Naehrstoffdatenbank. Karlsruhe. Lizenz CC BY 4.0. https://www.blsdb.de

Idempotent: baut die Datenbank jedes Mal neu und ersetzt sie atomar.

    docker compose --profile import run --rm import python -m app.import_bls
    python -m app.import_bls --file BLS_4_0_Daten_2025_DE.xlsx   # lokale Datei
"""

from __future__ import annotations

import argparse
import io
import logging
import os
import re
import zipfile

import requests
from openpyxl import load_workbook

from . import fooddb

log = logging.getLogger("import_bls")

BLS_URL = os.getenv("BLS_URL", "https://www.blsdb.de/assets/uploads/BLS_4_0_2025_DE.zip")
USER_AGENT = "kalorien-tracker/1.0 (privat)"

# Spalten-Praefix im BLS -> eigenes Feld
COLUMNS = {
    "ENERCC": "kcal", "PROT625": "protein", "CHO": "carbs", "SUGAR": "sugar",
    "FAT": "fat", "FASAT": "satfat", "FIBT": "fiber", "NACL": "salt",
}

# Typische Stueckgewichte (essbarer Anteil). Der BLS hat keine Portionsangaben.
PIECES = [
    (r"^Apfel (roh|geschält, roh)$", 150, "Apfel"),
    (r"^Banane roh$", 120, "Banane"),
    (r"^Birne roh$", 160, "Birne"),
    (r"^Orange roh|^Apfelsine", 180, "Orange"),
    (r"^Mandarine roh", 60, "Mandarine"),
    (r"^Kiwi", 75, "Kiwi"),
    (r"^Pfirsich roh", 130, "Pfirsich"),
    (r"^Tomate roh", 80, "Tomate"),
    (r"^Hühnerei roh$", 60, "Ei (M)"),
    (r"^Hühnerei .*gekocht", 60, "Ei (M)"),
    (r"^Hühnerei Eiklar, roh$", 35, "Eiklar"),
    (r"^Hühnerei Eigelb, roh$", 18, "Eigelb"),
    (r"^Weizenbrötchen$|^Mehrkornbrötchen$|^Roggenbrötchen", 60, "Brötchen"),
    (r"^Brezel|^Laugenbrezel", 70, "Brezel"),
    (r"^Knäckebrot", 10, "Scheibe"),
    (r"toast", 25, "Scheibe"),
    (r"brot\b", 50, "Scheibe"),
    (r"^Kartoffel (geschält|ungeschält), (roh|gekocht)", 90, "Kartoffel"),
    (r"^Karotte roh|^Möhre roh", 70, "Karotte"),
    (r"^Paprika(schote)? .*roh", 150, "Paprika"),
    (r"^Gurke roh|^Salatgurke roh", 400, "Gurke"),
    (r"^Avocado roh", 140, "Avocado"),
    (r"^Zwiebel roh", 60, "Zwiebel"),
    (r"^Walnuss", 5, "Walnuss"),
    (r"^Dattel", 8, "Dattel"),
]


LIQUID = re.compile(r"^(H-)?(Voll|Mager|Butter|Rohmilch/)?milch\b(?!.*pulver)|^Milch |^Kefir|^Ayran|drink\b|saft\b|nektar\b|schorle\b|^Sojadrink|^Kakao \(Getränk", re.I)


def _float(v) -> float | None:
    if v is None or v == "" or v == "-":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def _piece(name: str) -> tuple[float | None, str | None]:
    for pat, g, label in PIECES:
        if re.search(pat, name, re.I):
            return g, label
    return None, None


def load_workbook_bytes(file: str | None) -> bytes:
    if file:
        with open(file, "rb") as f:
            data = f.read()
    else:
        log.info("Lade %s", BLS_URL)
        r = requests.get(BLS_URL, headers={"User-Agent": USER_AGENT}, timeout=120)
        r.raise_for_status()
        data = r.content
    if data[:2] == b"PK" and (file is None or file.endswith(".zip")):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            name = next(n for n in z.namelist() if re.search(r"Daten.*\.xlsx$", n))
            log.info("Entpacke %s", name)
            return z.read(name)
    return data


def run(file: str | None = None) -> int:
    wb = load_workbook(io.BytesIO(load_workbook_bytes(file)), read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {}
    for i, h in enumerate(header):
        if not h:
            continue
        code = str(h).split(" ", 1)[0]
        if code in COLUMNS and "[" in str(h):  # Wertspalte, nicht "Datenherkunft"
            idx[COLUMNS[code]] = i
    missing = set(COLUMNS.values()) - set(idx)
    if missing:
        raise SystemExit(f"BLS-Format unbekannt, Spalten fehlen: {missing}")

    b = fooddb.Builder("bls")
    try:
        batch = []
        for r in rows:
            code, name = r[0], r[1]
            if not code or not name:
                continue
            item = {k: _float(r[i]) for k, i in idx.items()}
            if item["kcal"] is None:
                continue
            pg, pl = _piece(name)
            item.update(
                code=str(code), name=str(name).strip(), piece_g=pg, piece_label=pl,
                # N = alkoholfreie Getraenke, P = alkoholische Getraenke, dazu Milch, Drinks, Saefte
                liquid=1 if str(code)[0] in "NP" or LIQUID.search(str(name)) else 0,
                popularity=0,
            )
            batch.append(item)
        b.add_many(batch)
        n = b.finish({"source": "BLS 4.0", "license": "CC BY 4.0"})
    except BaseException:
        b.abort()
        raise
    log.info("BLS: %d Lebensmittel importiert", n)
    return n


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    ap = argparse.ArgumentParser(description="BLS 4.0 importieren")
    ap.add_argument("--file", help="lokale .zip oder .xlsx statt Download")
    args = ap.parse_args()
    run(args.file)


if __name__ == "__main__":
    main()
